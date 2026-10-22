"""Build the AeroGard research paper (IEEE two-column) using the example paper as the style
template: its styles, page setup and section layout are reused; all content is replaced.

    python scripts/build_paper.py [path/to/template.docx]
"""

import copy
import json
import math
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

REPO = Path(__file__).resolve().parents[1]
TEMPLATE = Path(sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\admin\Downloads\CropDiseaseDetection_ResearchPaper.docx")
OUT = REPO / "docs" / "paper" / "AeroGard_ResearchPaper.docx"
FIG = REPO / "docs" / "paper"
EXP = REPO / "docs" / "experiments"

doc = Document(str(TEMPLATE))
body = doc.element.body

# Keep the template's single-column section-break paragraph; drop all other content.
sect_break_proto = next(
    copy.deepcopy(p) for p in body.iterchildren(qn("w:p")) if p.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None
)
for child in list(body):
    if child.tag != qn("w:sectPr"):
        body.remove(child)

# --- inline markup: **bold**, *italic*, _{sub}, ^{sup} ---------------------------------------
TOKEN = re.compile(r"(\*\*.+?\*\*|\*.+?\*|_\{.+?\}|\^\{.+?\})")


def add_runs(par, text, size=None, bold=False, italic=False, small_caps=False):
    for part in TOKEN.split(text):
        if not part:
            continue
        b, i, sub, sup = bold, italic, False, False
        if part.startswith("**"):
            part, b = part[2:-2], True
        elif part.startswith("*"):
            part, i = part[1:-1], True
        elif part.startswith("_{"):
            part, sub = part[2:-1], True
        elif part.startswith("^{"):
            part, sup = part[2:-1], True
        r = par.add_run(part)
        r.bold, r.italic = b or None, i or None
        # Both map to w:vertAlign, so only ever set the one that applies.
        if sub:
            r.font.subscript = True
        elif sup:
            r.font.superscript = True
        if small_caps:
            r.font.small_caps = True
        if size:
            r.font.size = Pt(size)
    return par


def para(text="", *, size=None, align="both", first_line=True, before=None, after=60, bold=False, italic=False, keep_next=False, small_caps=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = {"both": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    if first_line:
        pf.first_line_indent = Pt(14.4)  # 288 twips, as in the template
    if before is not None:
        pf.space_before = Pt(before / 20)
    if after is not None:
        pf.space_after = Pt(after / 20)
    if keep_next:
        pf.keep_with_next = True
    add_runs(p, text, size=size, bold=bold, italic=italic, small_caps=small_caps)
    return p


def lead(label, text):
    """Paragraph starting with an italic run-in label, e.g. 'Security. ...'."""
    p = para("")
    add_runs(p, f"*{label}* {text}")
    return p


def heading(text):
    return doc.add_paragraph(text, style="Heading 1")


def bullet_free_list(items):
    for it in items:
        para(it)


# --- tables ---------------------------------------------------------------------------------
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]
_table_no = 0
_fig_no = 0


def _set(el, tag, **attrs):
    node = el.find(qn(tag))
    if node is None:
        node = OxmlElement(tag)
        el.append(node)
    for k, v in attrs.items():
        node.set(qn(f"w:{k}"), str(v))
    return node


def table(caption, rows, widths, center_cols=()):
    """IEEE-style table: 'TABLE N' + small-caps title, grey borders, shaded header row."""
    global _table_no
    _table_no += 1
    para(f"TABLE {ROMAN[_table_no - 1]}", size=8, align="center", first_line=False, before=120, after=0, keep_next=True)
    para(caption, size=8, align="center", first_line=False, after=60, keep_next=True, small_caps=True)
    t = doc.add_table(rows=len(rows), cols=len(widths))
    tbl = t._tbl
    tblPr = tbl.tblPr
    _set(tblPr, "w:tblW", w=sum(widths), type="dxa")
    _set(tblPr, "w:tblLayout", type="fixed")
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths, strict=True):
        gc.set(qn("w:w"), str(w))
    for r_i, (row, tr) in enumerate(zip(rows, t.rows, strict=True)):
        trPr = tr._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        if r_i == 0:
            trPr.append(OxmlElement("w:tblHeader"))
        for c_i, (val, cell) in enumerate(zip(row, tr.cells, strict=True)):
            tcPr = cell._tc.get_or_add_tcPr()
            _set(tcPr, "w:tcW", w=widths[c_i], type="dxa")
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                b = OxmlElement(f"w:{side}")
                for k, v in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "808080")):
                    b.set(qn(f"w:{k}"), v)
                borders.append(b)
            tcPr.append(borders)
            if r_i == 0:
                shd = OxmlElement("w:shd")
                for k, v in (("val", "clear"), ("color", "auto"), ("fill", "E7E6E6")):
                    shd.set(qn(f"w:{k}"), v)
                tcPr.append(shd)
            mar = OxmlElement("w:tcMar")
            for side, w in (("top", 20), ("left", 50), ("bottom", 20), ("right", 50)):
                m = OxmlElement(f"w:{side}")
                m.set(qn("w:w"), str(w))
                m.set(qn("w:type"), "dxa")
                mar.append(m)
            tcPr.append(mar)
            _set(tcPr, "w:vAlign", val="center")
            p = cell.paragraphs[0]
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if (r_i == 0 or c_i in center_cols) else WD_ALIGN_PARAGRAPH.LEFT
            add_runs(p, str(val), size=7.5, bold=(r_i == 0))
    para("", first_line=False, after=0)
    return t


def figure(path, caption, width_in=3.24):
    global _fig_no
    _fig_no += 1
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width_in))
    para(f"Fig. {_fig_no}. {caption}", size=8, first_line=False, before=40, after=120)


def equation(expr, number):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    for val, pos in (("center", 2520), ("right", 5040)):
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), val)
        tab.set(qn("w:pos"), str(pos))
        tabs.append(tab)
    pPr.append(tabs)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.add_run("\t")
    add_runs(p, expr, italic=True)
    p.add_run(f"\t({number})")


def reference(n, text):
    p = para("", size=8, first_line=False, after=20)
    pf = p.paragraph_format
    pf.left_indent = Inches(0.3)
    pf.first_line_indent = Inches(-0.3)
    pf.tab_stops.add_tab_stop(Inches(0.3))
    add_runs(p, f"[{n}]\t{text}", size=8)


# --- data -----------------------------------------------------------------------------------
exp = {p: json.loads((EXP / f"aq_forecast_{p}_1.1.0.json").read_text()) for p in ("pm25", "pm10", "no2", "o3")}
pm = exp["pm25"]
rows_split = pm["rows"]
lgb = {p: next(r for r in e["comparison_test"] if r["model"] == "lightgbm_direct") for p, e in exp.items()}
per_h = {m: pm["test"][m]["per_horizon"] for m in ("lightgbm_direct", "seasonal_naive", "persistence")}
val = {m: pm["validation"][m]["overall"]["mae"] for m in pm["validation"]}
best_iter = sorted(pm["hyperparameters"]["best_iterations"].values())
sn_mae = next(r for r in pm["comparison_test"] if r["model"] == "seasonal_naive")["mae"]
all_obs = {r["model"]: r["mae"] for r in pm["comparison_test_all_observed"]}
LBL = {"pm25": "PM_{2.5}", "pm10": "PM_{10}", "no2": "NO_{2}", "o3": "O_{3}", "so2": "SO_{2}", "co": "CO"}
def signed(v: float, plus: bool = False) -> str:
    """Typographic minus sign for negative numbers."""
    txt = f"{v:+.2f}" if plus else f"{v:.2f}"
    return txt.replace("-", "−")


MODEL = {"lightgbm_direct": "LightGBM", "seasonal_naive": "Seasonal naive", "persistence": "Persistence"}

risk = json.loads((EXP / "risk_1.0.0.json").read_text())
shap = json.loads((EXP / "risk_1.0.0_shap.json").read_text(encoding="utf-8"))
coh = risk["cohort"]
cnt = risk["splits"]["counts"]
rm = risk["models"]
served = rm["lightgbm_shared_heads"]["test_seen"]
unseen = rm["lightgbm_shared_heads"]["test_unseen"]
no_health = rm["ablation_no_health"]["test_seen"]
pers = risk["personalization"]["test_seen"]
shap_f = shap["parts"]["test_seen"]["all"]["factors"]
EXPOSURE_FACTORS = ("pm25_exposure", "pm25_change", "pm10_exposure", "o3_exposure", "exposure_completeness")
shap_exposure = sum(shap_f[k]["share"] for k in EXPOSURE_FACTORS)
shap_exposure_val = sum(shap["parts"]["validation"]["all"]["factors"][k]["share"] for k in EXPOSURE_FACTORS)
shap_diary = shap_f["symptoms"]["share"] + shap_f["flare_ups"]["share"]
shap_err_exp = math.ceil(math.log10(shap["parts"]["test_seen"]["additivity_max_abs_error"]))
RISK_MODEL = {
    "stratum_prior": "Stratum prior (no features)",
    "logistic": "Logistic regression",
    "lightgbm_pooled": "LightGBM pooled",
    "lightgbm_separate": "LightGBM per disease",
    "lightgbm_shared_heads": "**Shared trunk + heads (served)**",
    "ablation_no_exposure": "Ablation: no exposure",
    "ablation_no_health": "Ablation: no health",
}


def ci(m: dict, key: str = "auroc") -> str:
    c = m.get(f"{key}_ci95")
    return f"{m[key]:.3f} [{c['lo']:.3f}, {c['hi']:.3f}]" if c else f"{m[key]:.3f}"

# ============================================================================================
# Title block (single column)
# ============================================================================================
title = para(
    "AeroGard: Hyperlocal Air-Quality Forecasting, Exposure Estimation and Explainable Respiratory-Risk Alerts for Asthma and COPD",
    size=24, align="center", first_line=False, after=240,
)

authors = doc.add_table(rows=1, cols=3)
tblPr = authors._tbl.tblPr
_set(tblPr, "w:tblW", w=10440, type="dxa")
for gc in authors._tbl.tblGrid.findall(qn("w:gridCol")):
    gc.set(qn("w:w"), "3480")
for i, cell in enumerate(authors.rows[0].cells, start=1):
    tcPr = cell._tc.get_or_add_tcPr()
    _set(tcPr, "w:tcW", w=3480, type="dxa")
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "none")
        borders.append(b)
    tcPr.append(borders)
    lines = [(f"Author Name {i}", 11, False), ("Roll No.: ____________", 9, True),
             ("Department of Computer Science and Engineering", 9, True), ("Presidency University", 9, True)]
    for j, (txt, sz, it) in enumerate(lines):
        p = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(txt)
        r.font.size = Pt(sz)
        r.italic = it or None

body.insert(len(body) - 1, sect_break_proto)  # end of single-column section

# ============================================================================================
# Abstract, index terms
# ============================================================================================
abstract = doc.add_paragraph()
abstract.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
abstract.paragraph_format.space_after = Pt(3)
add_runs(abstract, "Abstract—", size=9, bold=True, italic=True)
add_runs(
    abstract,
    "People with asthma or chronic obstructive pulmonary disease (COPD) usually learn about air pollution from a "
    "city-wide index that describes the present, not the air around them tomorrow, and says nothing about their own "
    "exposure or health. We built AeroGard, a software prototype that forecasts air quality 24 hours ahead at a "
    "user’s location, estimates personal exposure, and turns these with a symptom diary into a calibrated, "
    "explained personal risk and an alert. For Bengaluru we ingested 753,481 hourly observations from 14 regulatory "
    "monitors (February 2025 to September 2026) with reanalysis weather. A direct multi-horizon LightGBM forecaster "
    f"reached a 24-hour PM_{{2.5}} mean absolute error of {lgb['pm25']['mae']:.2f} µg/m³ on an untouched test period, "
    f"{lgb['pm25']['skill_vs_persistence']:.0%} better than persistence. Because no respiratory outcome data exist for "
    f"the city, the risk layer was developed on a clearly labelled synthetic cohort of {coh['patients']} patients "
    "simulated inside the real pollution field. A LightGBM model with a shared trunk and asthma and COPD heads, "
    f"Platt-calibrated, reached an AUROC of {served['auroc']:.3f} (95% CI {served['auroc_ci95']['lo']:.3f}–"
    f"{served['auroc_ci95']['hi']:.3f}) with an expected calibration error of {served['ece']:.4f}, and "
    f"{unseen['auroc']:.3f} for patients never seen in training. The symptom diary carried the signal; estimated "
    "exposure added no measurable discrimination. Exact TreeSHAP explanations, scaled to add up to the probability "
    "shown, and an idempotent alert engine are served through an authenticated, consent-aware API. These are "
    "synthetic-data results, not clinical validation.",
    size=9,
)
kw = doc.add_paragraph()
kw.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
kw.paragraph_format.space_after = Pt(6)
add_runs(kw, "Index Terms—", size=9, bold=True, italic=True)
add_runs(kw, "air quality forecasting, PM2.5, LightGBM, exposure assessment, asthma, COPD, risk prediction, probability calibration, SHAP, synthetic cohort, data leakage", size=9)

# ============================================================================================
heading("I. Introduction")
para(
    "Air pollution is one of the leading environmental risks to health worldwide [1]. In India it was linked to "
    "about 1.67 million deaths in 2019 [2]. Short-term rises in particulate matter, nitrogen dioxide and ozone are "
    "associated with higher mortality [3] and with more asthma emergency visits and hospital admissions [4], so "
    "people living with asthma or COPD are among those who most need timely and relevant warnings."
)
para(
    "In practice most people check a city-level Air Quality Index. That number describes the city now, not the "
    "neighbourhood tomorrow, and it ignores where a person spends their time and how sensitive they are. Building "
    "even a reliable picture of the present is harder than it looks: monitoring stations are sparse, public feeds "
    "arrive late, readings drop out, and, as we found, units can be mislabelled."
)
para(
    "Here we describe AeroGard, a system that forecasts air quality for each of the next 24 hours at a user’s "
    "location, estimates personal exposure from location and microenvironment, and converts this, together with a "
    "symptom diary, into a calibrated personal risk, an explanation and an alert. The contributions are: (1) a "
    "validated ingestion pipeline for Bengaluru’s regulatory monitors, including the detection and correction of a "
    "unit error in the public feed; (2) a leakage-audited 24-hour forecaster compared on equal terms with two "
    "baselines; (3) a transparent exposure estimate; (4) a synthetic asthma and COPD cohort simulated inside the "
    "real pollution field, on which a calibrated risk model is compared with baselines and ablations under a "
    "protocol fixed before testing; and (5) exact, additive explanations and an idempotent alert engine served "
    "through a secured API. We report negative findings as such: on this cohort, pollution exposure did not "
    "improve risk discrimination."
)

heading("II. Problem Statement and Objectives")
para(
    "Generic air-quality monitoring reports environmental conditions but does not answer the question a patient "
    "actually has: given the air expected around me, my recent exposure and my health, what is my respiratory risk "
    "over the next 24 hours? There is a need for a system that produces location-specific forecasts, estimates "
    "exposure honestly, and communicates risk without making clinical claims."
)
para(
    "Our objectives are to (1) collect and validate hourly pollutant and weather data for Bengaluru; (2) match a "
    "user’s location to nearby monitors; (3) forecast pollutant concentrations for each of the next 24 hours and "
    "evaluate the forecasts against observed data; (4) estimate personal exposure from location and activity "
    "context; (5) predict calibrated asthma and COPD risk and map it to alert levels; and (6) deliver the results "
    "through a secure API and an Android application. Objectives (1)–(5) and the API are complete; the Android "
    "application is in progress."
)

heading("III. Literature Review")
para("Table I lists ten studies we reviewed, with the method and data each used, its main finding and its main limitation.")
table(
    "Summary of Related Work",
    [
        ["Ref.", "Year", "Method / Data", "Key Finding", "Limitation"],
        ["[2]", "2021", "Burden-of-disease estimates for Indian states", "Air pollution linked to about 1.67 M deaths in India in 2019", "Population level; no individual risk"],
        ["[3]", "2020", "Meta-analysis of short-term exposure studies", "PM, NO_{2} and O_{3} associated with higher mortality", "Pooled effects; not a prediction tool"],
        ["[4]", "2015", "Meta-analysis of time-series studies", "Pollutants associated with asthma ER visits and admissions", "Aggregate outcomes, not personal"],
        ["[5]", "2012", "Review of real-time AQ forecasting", "Surveys chemical-transport and statistical approaches", "Predates current ML methods"],
        ["[6]", "2017", "LSTM with spatio-temporal inputs, Beijing", "Outperformed statistical baselines for PM_{2.5}", "One city; data-hungry model"],
        ["[7]", "2012", "Multi-step forecasting strategies, NN5 data", "Direct and multi-output strategies beat recursive", "Not specific to air quality"],
        ["[8]", "2004", "Interpolation of monitor data (IDW, kriging, others)", "Method choice changes exposure estimates", "Depends on monitor density"],
        ["[9]", "2011", "Review of indoor/outdoor particle relationships", "I/O ratios vary with building and ventilation", "Few South-Asian buildings"],
        ["[10]", "2022", "Connected inhaler, peak-flow and wearable data + ML", "Framework for predicting asthma attacks", "Protocol stage; no AQ forecast"],
        ["[11]", "2014", "Review of air pollution in Indian cities", "Documents sources and sparse monitoring", "Descriptive; no forecasting"],
    ],
    [420, 440, 1330, 1550, 1300],
    center_cols=(0, 1),
)
para(
    "Epidemiological work establishes that short-term exposure matters [1]–[4], but its results are population "
    "averages rather than personal predictions. Forecasting research ranges from chemical-transport models to "
    "data-driven methods [5]; deep sequence models can outperform statistical baselines [6], and the choice between "
    "recursive and direct multi-step strategies affects accuracy [7]. Exposure studies show that an outdoor monitor is "
    "an imperfect proxy for what a person breathes: the interpolation method [8] and indoor infiltration [9] both "
    "change the estimate. Connected-device asthma studies [10] predict attacks from physiological and inhaler data "
    "but do not use air-quality forecasts. Finally, results in this area are easy to inflate through data leakage, "
    "where information unavailable at prediction time enters the model [12]."
)
lead(
    "Research gap.",
    "Existing work treats air-quality forecasting, exposure estimation and respiratory risk separately. We found few "
    "systems for Indian cities that combine station-level 24-hour forecasts, a transparent exposure estimate and "
    "personal risk in one traceable pipeline, evaluated with explicit leakage checks against simple baselines.",
)

heading("IV. Proposed Methodology")
figure(FIG / "fig1_pipeline.png", "Methodology pipeline. The dashed stage is in progress.")
para(
    "Fig. 1 shows the pipeline. Hourly pollutant means from regulatory monitors are obtained through the OpenAQ "
    "API [13], and hourly ERA5 reanalysis weather [14] through Open-Meteo [15]. Records are validated and cleaned, "
    "matched to the user’s location, turned into features and passed to the 24-hour forecaster. The forecast and the "
    "user’s location history drive the exposure estimate, which, together with the health profile and symptom "
    "diary, feeds the respiratory-risk model, its calibration, the explanation and the alert engine."
)
lead(
    "Validation and cleaning.",
    "Every record is converted to a canonical form: a UTC timestamp at the end of the averaging hour, a pollutant "
    "identity, and a concentration in µg/m³. Gases reported in ppb are converted at 25 °C and 1 atm "
    "(µg/m³ = ppb × M / 24.45, with M the molecular weight). Records with missing timestamps, unknown pollutants, "
    "invalid coordinates, or missing, non-finite or negative values are rejected and the reason is counted. Hours "
    "with fewer than 75% of their sub-hourly readings, provider flags or physically implausible values are kept but "
    "flagged. Short gaps of up to 3 h are forward-filled from past values only, and outliers are flagged with a "
    "trailing robust z-score (median and MAD over the previous 24 h) rather than deleted.",
)
lead(
    "Hyperlocal matching.",
    "For a location we find reference monitors within 10 km by haversine distance, discard readings older than "
    "3 h or later than the query time, and combine up to four stations by inverse-distance weighting with power 2 "
    "[8], halving the weight of low-coverage readings. If no usable station remains, the system reports the value as "
    "unavailable instead of extrapolating.",
)

heading("V. Dataset")
para(
    "The study area is Bengaluru, India. A 25 km search around the city centre returned 14 CPCB/KSPCB reference "
    "monitors with data in the study period; low-cost sensors listed on the same platform were excluded from ground "
    "truth. All stations share a gap from about October 2022 to February 2025, when the source feed changed, so we "
    "use 18 February 2025 to 21 September 2026, about 19 months. Hourly weather (13,992 h) comes from ERA5 at the city "
    "centre. Table II summarises the pollutant data; about 0.3% of fetched records were rejected, all for missing values."
)
from sqlalchemy import text  # noqa: E402

sys.path.insert(0, str(REPO / "backend"))
from app.db.session import get_sessionmaker  # noqa: E402

with get_sessionmaker()() as db:
    prow = db.execute(text("""
        select pollutant, count(distinct station_id), count(*),
               round(100.0*count(*) filter (where quality_flag<>'valid')/count(*),1),
               round(percentile_cont(0.5) within group (order by concentration)::numeric,1)
        from air_quality_observations group by pollutant order by count(*) desc""")).all()
table(
    "Dataset Composition (Bengaluru, Feb 2025 – Sep 2026)",
    [["Pollutant", "Stations", "Hourly rows", "Flagged", "Median (µg/m³)"]]
    + [[LBL[p], s, f"{n:,}", f"{fl}%", f"{med}"] for p, s, n, fl, med in prow],
    [1000, 900, 1100, 900, 1140],
    center_cols=(1, 2, 3, 4),
)
para(
    "The current feed labels CO, NO_{2} and SO_{2} in ppb. For CO this cannot be right: at BTM Layout the median "
    "reading was 0.49 in June–September 2025 against 1,260 µg/m³ on the older feed in the same months of 2021, and "
    "ambient CO is typically several hundred ppb. The values match CPCB’s reporting unit for CO, mg/m³ [16], so all "
    "CO rows were converted accordingly. For NO_{2} and SO_{2} the comparison is ambiguous; their absolute levels are "
    "therefore treated as unverified and are not used in alert thresholds. This does not affect forecasting, because "
    "tree-based models are unchanged by a constant rescaling of an input."
)

heading("VI. Model Development")
para(
    "Each training row is a (station, origin hour t) pair with 35 features: the current value of the target "
    "pollutant and whether it was filled; lags of 1, 6, 12, 24 and 48 h; the 1-h difference; rolling mean, maximum "
    "and standard deviation over 6 h and 24 h; the number of observed hours in the last day; the current values of "
    "the other pollutants; weather at t (temperature, humidity, wind speed and direction as sine and cosine, "
    "pressure, precipitation and its 24-h total); calendar terms in Indian Standard Time; and station coordinates. "
    "Every feature declares how far back it reads, and an automated test confirms that altering any data after t "
    "leaves the features at t unchanged, guarding against leakage [12]."
)
para(
    "We use a direct multi-horizon strategy [7]: one LightGBM gradient-boosted tree model [17] per horizon "
    "h = 1, …, 24, each trained on log(1 + y) and back-transformed, with negative outputs clipped to zero. No forecast "
    "is ever fed back in as an input. We chose LightGBM because it trains quickly on a CPU, handles missing values "
    "natively, performs well on tabular data and produces small artifacts that are easy to deploy. Two baselines "
    "complete the comparison: persistence, ŷ(t+h) = y(t), and seasonal naive, ŷ(t+h) = y(t+h−24)."
)
para(
    "Targets are observed values from hours flagged valid. Inputs may still use flagged hours, but isolated PM_{2.5} "
    "readings of 600–900 µg/m³ during the monsoon, mostly at 25–50% coverage, looked like sensor faults and were "
    "not used as ground truth. Results on all observed hours are reported as a secondary check."
)

heading("VII. Experimental Setup")
para("Table III gives the training configuration and the environment used for the experiments.")
table(
    "Experimental Setup",
    [
        ["Item", "Value"],
        ["Split (chronological)", "Train 18 Feb 2025 – 31 Jan 2026; validation 1 Feb – 30 Apr 2026; test 1 May – 21 Sep 2026"],
        ["Origins (PM_{2.5})", f"Train {rows_split['train']:,}; validation {rows_split['validation']:,}; test {rows_split['test']:,}"],
        ["Embargo", f"{pm['embargo_hours']} h before each boundary, so no target crosses a split"],
        ["Model", "LightGBM, 24 direct models per pollutant"],
        ["Hyperparameters", "≤ 600 trees, learning rate 0.03, 31 leaves, min. 50 samples per leaf, row and column subsampling 0.8, L2 = 1"],
        ["Early stopping", f"50 rounds on validation (best {best_iter[0]}–{best_iter[-1]} trees)"],
        ["Target transform", "log(1 + y)"],
        ["Evaluation", "Test period used once; all models scored on identical pairs"],
        ["Software", "Python 3.14, LightGBM 4.7, scikit-learn 1.9, pandas 3.0, PostgreSQL 16"],
        ["Hardware", "Consumer laptop CPU, Windows 11; no GPU"],
    ],
    [1300, 3740],
)

heading("VIII. Evaluation Metrics")
para(
    "We report mean absolute error (MAE), root mean squared error (RMSE), the coefficient of determination (R²), "
    "mean bias, and skill relative to persistence. With ŷ_{i} the forecast, y_{i} the observation and n the number "
    "of forecast–observation pairs,"
)
equation("MAE = (1/n) Σ |ŷ_{i} − y_{i}|", 1)
equation("RMSE = √[(1/n) Σ (ŷ_{i} − y_{i})²]", 2)
equation("R² = 1 − Σ (ŷ_{i} − y_{i})² / Σ (y_{i} − ȳ)²", 3)
equation("Skill = 1 − MAE_{model} / MAE_{persistence}", 4)
para(
    "where ȳ is the mean observation. Metrics are computed per horizon and pooled over all 24 horizons. To keep "
    "the comparison fair, every model is scored on the same (origin, horizon) pairs, namely those where all models "
    "produce a forecast and a valid observation exists; otherwise a baseline that cannot forecast difficult cases "
    "would look better than it is."
)

heading("IX. Forecast Results")
para(
    "Table IV gives pooled test-set performance for four pollutants and Table V the PM_{2.5} error by horizon. The "
    "test period covers the monsoon season."
)
res_rows = [["Pollutant", "Model", "MAE (µg/m³)", "RMSE (µg/m³)", "R²", "Skill"]]
for p in ("pm25", "pm10", "no2", "o3"):
    for r in exp[p]["comparison_test"]:
        res_rows.append([
            LBL[p], MODEL[r["model"]], f"{r['mae']:.2f}", f"{r['rmse']:.2f}", signed(r["r2"]),
            f"{r['skill_vs_persistence']:+.0%}",
        ])
table("Test Set Performance (Valid Hours)", res_rows, [850, 1270, 700, 750, 700, 770], center_cols=(2, 3, 4, 5))
hs = (1, 3, 6, 12, 18, 24)
table(
    "PM_{2.5} MAE by Forecast Horizon",
    [["Model (MAE, µg/m³)"] + [f"{h} h" for h in hs]]
    + [[MODEL[m]] + [f"{per_h[m][f'h{h}']['mae']:.2f}" for h in hs] for m in ("lightgbm_direct", "seasonal_naive", "persistence")],
    [1370, 610, 610, 610, 610, 610, 620],
    center_cols=(1, 2, 3, 4, 5, 6),
)
para(
    f"LightGBM had the lowest MAE for PM_{{2.5}}, PM_{{10}} and NO_{{2}}, improving on persistence by "
    f"{lgb['pm25']['skill_vs_persistence']:.0%}, {lgb['pm10']['skill_vs_persistence']:.0%} and "
    f"{lgb['no2']['skill_vs_persistence']:.0%}. Table V shows where the gain comes from: at 1 h ahead persistence is "
    "almost as good, but its error grows quickly with horizon, whereas LightGBM’s grows slowly and stays below the "
    "seasonal baseline at every horizon. The same ordering held on the validation period (February–April, the dry "
    f"season, when PM_{{2.5}} is higher): MAE {val['lightgbm_direct']:.2f} for LightGBM against "
    f"{val['seasonal_naive']:.2f} for seasonal naive and {val['persistence']:.2f} for persistence."
)
para(
    "Three results qualify this picture. For O_{3}, the seasonal-naive baseline has the lowest MAE, although LightGBM "
    "has the lower RMSE and higher R²; ozone follows the daily cycle of sunlight so closely that yesterday’s value at "
    "the same hour is hard to beat. For NO_{2} at 1 h ahead, persistence beats LightGBM. For PM_{2.5}, RMSE is well "
    "above MAE and R² is modest, which means the model improves typical hours but does not anticipate sudden spikes. "
    f"All forecasts are slightly biased low (PM_{{2.5}} bias {signed(lgb['pm25']['bias'], plus=True)} µg/m³), as expected from "
    "training on a log scale. Scoring against all observed hours, including flagged ones, gives the same conclusion "
    f"(PM_{{2.5}} MAE {all_obs['lightgbm_direct']:.2f} against {all_obs['persistence']:.2f} for persistence)."
)

heading("X. Exposure Estimation")
para(
    "Exposure is an estimate, not a measurement: no personal sensor is used. For each hour we take the ambient "
    "concentration at the user’s location from the matching in Section IV and multiply it by a factor for the "
    "user’s microenvironment (Table VI). The factors are central assumptions drawn from indoor/outdoor studies [9], "
    "not values calibrated to AeroGard users."
)
table(
    "Microenvironment Factors (Exposure / Ambient)",
    [
        ["Setting", "PM_{2.5}", "PM_{10}", "NO_{2}", "O_{3}", "CO"],
        ["Outdoors / unknown", "1.0", "1.0", "1.0", "1.0", "1.0"],
        ["Indoors, home", "0.8", "0.6", "0.6", "0.3", "0.9"],
        ["Indoors, other", "0.6", "0.5", "0.5", "0.2", "0.8"],
        ["In traffic", "1.3", "1.3", "1.6", "0.8", "1.8"],
    ],
    [1540, 700, 700, 700, 700, 700],
    center_cols=(1, 2, 3, 4, 5),
)
para(
    "If no location fix exists in the two hours before a given hour, that hour’s exposure is reported as missing "
    "rather than guessed, and an unknown microenvironment is treated as outdoors and flagged. Six-, 24- and 72-hour "
    "means, maxima and cumulative exposure (µg/m³·h) are reported together with their coverage and marked incomplete "
    "below 75%. Because the factors are assumptions, the risk results were re-run with the factors scaled and with "
    "the exact simulated exposure (Section XIV)."
)

heading("XI. System Architecture and Implementation")
figure(FIG / "fig2_architecture.png", "System architecture of the platform.")
para(
    "Fig. 2 shows the architecture. Clients call a versioned REST API built with FastAPI. Routes stay thin and "
    "delegate to services, which use repositories for database access and the domain and machine-learning modules "
    "for matching, cleaning, features and inference. An ingestion command-line tool pulls data from the external "
    "sources, validates it and writes it to PostgreSQL."
)
lead(
    "Backend.",
    "The API exposes current, historical and forecast air quality for a location, reference stations, location "
    "upload, exposure, symptom-diary records, the current and 24-hour rolling risk with its explanation, risk "
    "history, and alerts with acknowledgement. Models come from frozen, versioned files; each carries a "
    "metadata record with its training period, metrics and checksum, and the service refuses to load a file whose "
    "checksum or feature list does not match the code. If no model is available the API returns an explicit error "
    "rather than a placeholder value.",
)
lead(
    "Security and privacy.",
    "Passwords are hashed with Argon2 and sessions use short-lived JSON Web Tokens with rotating refresh tokens. "
    "Every patient-scoped operation checks ownership, so one user cannot read another user’s data, and location and "
    "exposure require recorded consent. Sensitive actions are written to an audit log that stores counts and field "
    "names, never health or location values.",
)
lead(
    "Testing.",
    "An automated suite of 401 tests runs against a separate database built through the real migrations. It covers "
    "authentication and token misuse, patient isolation, source adapters under malformed input, timeouts and partial "
    "failure, validation and unit conversion, idempotent reloading, spatial edge cases, leakage guards for both "
    "feature sets, forecast and risk horizons, probability bounds, calibration, explanation additivity and "
    "consistency, alert rules and deduplication, and artifact checks.",
)

heading("XII. Synthetic Respiratory Cohort")
para(
    "No respiratory outcome data exist for Bengaluru, and collecting them requires ethics approval and a follow-up "
    "far longer than this project. The risk layer was therefore developed on a simulated panel of patients. The "
    "air-quality and weather inputs are real; the patients, their routines, exposure, symptoms and exacerbations "
    "are simulated. Every risk result in this paper is a synthetic-data evaluation, not clinical validation."
)
para(
    f"The cohort has {coh['patients']} patients ({coh['asthma_patients']} asthma, {coh['copd_patients']} COPD) with "
    "severity, age, home and workplace drawn from configured distributions, hidden frailty and pollution "
    "sensitivity, and daily routines (home, commute, work, errands, walks). True personal exposure is the real "
    "station field interpolated to each place, perturbed for street-scale variation and scaled by a personal "
    "microenvironment factor. The system only sees what production would see: location logged with a patient-"
    "specific rate and day-long dropouts, an unknown setting 10% of the time, and exposure estimated by the live "
    "matching rule. Events follow an hourly hazard,"
)
equation("log λ = c_{s} + u + sβΔ_{2.5} + β^{T}z + β_{I}I + β_{C}C", 5)
para(
    "for patient i at hour t (indices omitted), where c_{s} is solved so each disease × severity stratum has its "
    "configured annual rate, u is frailty, s pollution sensitivity, Δ_{2.5} the 24-h exponential moving average of true PM_{2.5} exposure above "
    "25 µg/m³ per 10 µg/m³, z collects PM_{10}, O_{3}, cold, humidity, recent-event and nocturnal-asthma terms, "
    "I is a seasonal viral infection, the dominant trigger of exacerbations [23], and C is a latent daily "
    "loss-of-control state that also drives the diary, standing for the worsening that precedes most "
    "exacerbations. A diary at 21:00 reports a 0–10 symptom score, reliever puffs and an exacerbation flag, filled "
    "with patient-specific adherence; 40% of patients also wear a device. The pollution effects are of the order "
    "reported in short-term exposure studies [3], [4] but were chosen to give a learnable signal and are not "
    "calibrated to Indian epidemiology. A scale factor on all pollution terms provides a negative control (zero) "
    "and stronger scenarios."
)
para(
    f"The simulation produced {coh['events']:,} events ({coh['events_per_patient_year']:.1f} per patient-year), "
    f"{coh['diary_symptom_completeness']:.0%} diary completeness and {coh['location_availability']:.0%} location "
    "availability. The target is whether an event starts within the next 24 hours of an origin; two random origins "
    "are drawn per patient-day, excluding warm-up, the final day and ongoing episodes."
)

heading("XIII. Risk Model, Calibration, Explanations and Alerts")
para(
    "Each origin has 38 features computed only from data at or before it, with the same code in training and in "
    "the API: estimated exposure (PM_{2.5} now, 6-, 24- and 72-h means, 24-h maximum, change and coverage; PM_{10} "
    "and O_{3} summaries), weather, calendar terms, profile (age, sex, disease, severity) and health history "
    "(last and mean symptom scores, trend, reliever use, diary days, events in 90 days, days since the last event, "
    "and the last wearable readings). Missing values stay missing, as in the training data, so the model learns a "
    "branch for them. A test perturbs every input after the origin and requires identical features."
)
para(
    "The model is a shared LightGBM trunk trained on all patients with disease as a feature, plus one small "
    "LightGBM head per disease boosted from the trunk’s output,"
)
equation("p = σ(m_{trunk}(x) + m_{d}(x))", 6)
para(
    "where σ is the logistic function and early stopping on validation keeps a head near zero if the disease adds "
    "nothing. Baselines use the same protocol: the base rate of the patient’s stratum, logistic regression with "
    "missing-value indicators, one pooled LightGBM, and one LightGBM per disease; two ablations remove the exposure "
    f"and the health features. Of the seen patients (80%), origins up to 31 January 2026 train "
    f"({cnt['train']['rows']:,} rows, {cnt['train']['events']:,} events), February–April validate "
    f"({cnt['validation']['rows']:,}, {cnt['validation']['events']}) and May–September test "
    f"({cnt['test_seen']['rows']:,}, {cnt['test_seen']['events']:,}); the other 20% of patients are tested over the "
    f"same months ({cnt['test_unseen']['rows']:,}, {cnt['test_unseen']['events']}). Origins whose label window "
    "crosses a boundary are dropped, the protocol was fixed before testing, and 95% intervals come from a "
    "patient-level cluster bootstrap with 200 resamples."
)
para("Probabilities are calibrated per disease on validation with Platt scaling [19],")
equation("p_{cal} = σ(a · logit p + b)", 7)
para(
    "after isotonic regression was rejected because its lowest block reported an exact zero risk for 4–6% of test "
    "predictions. Discrimination is measured by AUROC and AUPRC, accuracy of the probabilities by the Brier score "
    "[20] and the expected calibration error over B = 10 equal-frequency bins [21],"
)
equation("ECE = Σ_{b} (n_{b}/n) |ō_{b} − p̄_{b}|", 8)
para(
    "where ō_{b} and p̄_{b} are the observed rate and mean prediction in bin b. A versioned policy maps the "
    "calibrated probability to Low, Moderate, High or Severe per disease, with thresholds fixed on validation: "
    "Moderate at 1.5 × the base rate, High at the F_{1}-optimal alert threshold, and Severe at the top 1% of "
    "validation predictions. An optional personal offset shifts the log-odds by the patient’s own history,"
)
equation("p′ = σ(logit p_{cal} + δ)", 9)
para(
    "with δ = clip(ln((O + 2)/(E + 2)), ±ln 2), O the exacerbations the patient reported before the prediction "
    "time and E the number expected from their "
    "stratum over the same diary days, so the offset shrinks to zero for new users."
)
lead(
    "Explanations.",
    "Each prediction is explained with exact TreeSHAP [18], [22] computed by LightGBM. SHAP values are additive, "
    "so the trunk and head attributions sum, and because Platt scaling is linear in log-odds the attributions can "
    "be moved onto the probability the patient sees:",
)
equation("logit p′ = (aφ_{0} + b) + Σ_{j} aφ_{j}(x) + δ", 10)
para(
    "where φ_{0} is the base value and φ_{j} the attribution of feature j. The 38 features are grouped into 14 "
    "patient-facing factors, and the largest five are worded as having raised or lowered the estimate, never as "
    "causes. Each explanation is stored with the prediction’s model, feature and calibration versions and a "
    "fingerprint of its input vector, and is refused on read if any of them differs."
)
lead(
    "Alerts.",
    "An alert is created when the current estimate is valid, at least High, made no more than 2 h before "
    "evaluation, and no alert exists for that prediction or for an equivalent one, meaning a non-cancelled alert at "
    "the same or a higher level whose 24-h window is still running. A higher level supersedes open lower alerts; a "
    "recomputed prediction that falls below High withdraws its own alert. Predictions and the alert decision are "
    "written in one transaction with a row lock on the patient. An alert states only the level and percentage, "
    "never medication advice.",
)

heading("XIV. Risk Results")
para(
    f"Table VII compares the models on the test period. The test prevalence is {served['prevalence']:.1%} per "
    "24 hours, so a random model has an AUPRC of about 0.015."
)
risk_rows = [["Model", "AUROC [95% CI]", "AUPRC", "Brier", "ECE"]]
for key, label in RISK_MODEL.items():
    m = rm[key]["test_seen"]
    risk_rows.append([label, ci(m), f"{m['auprc']:.3f}", f"{m['brier']:.5f}", f"{m['ece']:.4f}"])
table("Risk Model Comparison, Test, Seen Patients (Synthetic Labels)", risk_rows, [1450, 1550, 640, 760, 640], center_cols=(1, 2, 3, 4))
para(
    f"The served model reached an AUROC of {served['auroc']:.3f} and generalised to unseen patients "
    f"({unseen['auroc']:.3f}). Its expected calibration error was {served['ece']:.4f}; Fig. 3 shows the observed "
    "event rate slightly above the prediction in most bins, because the monsoon test period has more infections "
    "than the February–April period on which the calibrators were fitted, a seasonal base-rate shift. Three "
    "findings should be stated plainly. First, the symptom diary carries the predictive signal: without health "
    f"features the AUROC falls to {no_health['auroc']:.3f}, close to the stratum prior. Second, the shared-head "
    "architecture is not measurably better than one pooled LightGBM or logistic regression; all intervals overlap. "
    "Third, removing the exposure features does not lower discrimination, at the default effect size or at "
    "two and four times it."
)
figure(FIG / "fig3_reliability.png", "Reliability of the calibrated risk, test, seen patients (10 equal-frequency bins).")
para(
    "Two reasons explain the third finding. The test period is the monsoon, when exposure is low and narrow "
    "(PM_{2.5} averaging 17 µg/m³, mostly below the 25 µg/m³ reference of the hazard), so the pollution term varies "
    "little; and the simulated symptom score already responds to pollution, so the diary absorbs most of it. The "
    "pre-registered test period was not moved."
)
lv = risk["policy_levels_test_seen"]
table(
    "Risk Levels on Test: Share of Time and Observed Event Rate",
    [["Disease", "Measure", "Low", "Moderate", "High", "Severe"]]
    + [
        [d.upper() if d == "copd" else d.capitalize(), meas]
        + [f"{100 * lv[d][lvl][key]:.1f}%" for lvl in ("Low", "Moderate", "High", "Severe")]
        for d in ("asthma", "copd")
        for meas, key in (("Share of time", "share"), ("Event rate", "event_rate"))
    ],
    [760, 1040, 760, 820, 760, 900],
    center_cols=(2, 3, 4, 5),
)
para(
    "Event rates rise with every level for both diseases (Table VIII), so the levels separate risk. Table IX "
    "compares alerting strategies, with High or above as an alert for the model."
)
PERS = [
    ("generic_aq_threshold", "Generic AQ threshold (NAQI Poor)"),
    ("disease_specific_aq_threshold", "Disease-specific AQ threshold"),
    ("model_global_threshold", "Model, one global threshold"),
    ("model_disease_threshold", "Model, per-disease thresholds"),
    ("model_adaptive_personal", "Model + personal offset"),
]
table(
    "Alerting Strategies, Test, Seen Patients",
    [["Strategy", "Precision", "Recall", "F_{1}", "Time alerting"]]
    + [
        [label, f"{pers[k]['precision']:.3f}", f"{pers[k]['recall']:.3f}", f"{pers[k]['f1']:.3f}", f"{pers[k]['alert_rate']:.1%}"]
        for k, label in PERS
    ],
    [1900, 780, 700, 640, 1020],
    center_cols=(1, 2, 3, 4),
)
para(
    "Bengaluru’s air almost never crosses the national index thresholds, so alerting on air quality alone is "
    "nearly silent. The personal offset raised recall from "
    f"{pers['model_disease_threshold']['recall']:.2f} to {pers['model_adaptive_personal']['recall']:.2f} at the cost "
    f"of alerting more often ({pers['model_disease_threshold']['alert_rate']:.1%} to "
    f"{pers['model_adaptive_personal']['alert_rate']:.1%} of the time) while leaving probability quality unchanged: "
    "a different operating point, not better discrimination. Even at High, most alerts are not followed by an "
    "event, which is expected at a 1.5% base rate and is why the level and percentage are both shown."
)
sens = risk["exposure_sensitivity"]
SENS = [
    ("estimated", "Estimated (production rule)"),
    ("factors_x0.8", "Microenvironment factors × 0.8"),
    ("factors_x1.2", "Microenvironment factors × 1.2"),
    ("home_ambient_only", "Ambient at home only"),
    ("true_exposure_oracle", "True exposure (oracle)"),
]
table(
    "Exposure Sensitivity: 24-h PM_{2.5} Exposure Error and Risk AUROC",
    [["Exposure used", "MAE (µg/m³)", "Bias", "Spearman ρ", "AUROC"]]
    + [
        [
            label,
            f"{sens[k]['exposure_error']['mae_24h_mean_pm25']:.2f}",
            signed(sens[k]["exposure_error"]["bias_24h_mean_pm25"], plus=True),
            f"{sens[k]['exposure_error']['spearman']:.3f}",
            f"{sens[k]['test_seen']['auroc']:.3f}",
        ]
        for k, label in SENS
    ],
    [1900, 820, 700, 820, 800],
    center_cols=(1, 2, 3, 4),
)
para(
    "Table X re-trains the model on each exposure variant. The production estimate is the least biased; ignoring "
    "mobility raises the error by about 44% and overstates exposure, because people spend time indoors. Even the "
    "exact simulated exposure does not change the AUROC, consistent with the finding above."
)
para(
    "Fig. 4 shows the global importance of the factors as the share of mean absolute SHAP value on the test "
    f"period. Recent symptom scores and flare-ups account for {shap_diary:.0%} of the attribution. Exposure factors "
    f"together account for {shap_exposure:.1%} ({shap_exposure_val:.1%} on validation), even though removing them "
    "does not lower the AUROC. The two results are consistent: SHAP describes what the trained model uses, and the "
    "ablation what it needs. Exposure is correlated with season and with the simulated symptom response, so a "
    "model trained without it reaches the same discrimination through those features. The additivity of equation "
    f"(10) held to within 10^{{{shap_err_exp}}} "
    "on every test prediction."
)
figure(FIG / "fig4_shap.png", "Global importance: share of mean |SHAP| per factor, test, seen patients. Lighter bars are exposure factors.")

heading("XV. Limitations")
para(
    "All risk results come from a simulated cohort whose data-generating process we wrote; they show that the "
    "pipeline, calibration, explanations and alerts behave correctly and honestly, not that AeroGard predicts real "
    "exacerbations. Its effect sizes are assumptions, so the finding that exposure adds nothing is a property of "
    "this cohort and test season, not evidence about pollution and asthma. The exposure factors are not validated "
    "with personal sensors, NO_{2} and SO_{2} units remain unverified, and only one city and one monsoon test "
    "season were evaluated. Clinical use would require real, ethically collected outcomes and re-derived "
    "thresholds."
)

heading("XVI. Conclusion")
para(
    "We built an end-to-end respiratory-alert system for Bengaluru, from validated ingestion to explained alerts. "
    "Careful validation mattered throughout: it exposed a CO unit error in the public feed, isolated suspect "
    "sensor spikes, caught a leakage bug in our own gap filling, and replaced a calibrator that reported "
    f"impossible zero risks. The direct LightGBM forecaster reduced 24-hour PM_{{2.5}} error to "
    f"{lgb['pm25']['mae']:.2f} µg/m³, {lgb['pm25']['skill_vs_persistence']:.0%} better than persistence. On a "
    f"clearly labelled synthetic cohort the calibrated risk model reached an AUROC of {served['auroc']:.2f} for "
    f"seen and {unseen['auroc']:.2f} for unseen patients, the symptom diary carried the signal, and exposure did "
    "not add measurable value. Explanations add up exactly to the probability shown, and alerts are raised once, "
    "traceably and without medication advice."
)

heading("XVII. Future Scope")
para(
    "Future work includes the Android application with push alerts and optional Health Connect data; a "
    "feasibility pilot with real patients under ethics approval; evaluation over a full year including winter, "
    "when exposure is higher; verification of NO_{2} and SO_{2} units against CPCB’s own feed; comparison with "
    "sequence models such as LSTM [6]; forecast-driven weather inputs; and validation of the exposure factors with "
    "personal sensors."
)

heading("References")
REFS = [
    "GBD 2019 Risk Factors Collaborators, “Global burden of 87 risk factors in 204 countries and territories, 1990–2019: A systematic analysis for the Global Burden of Disease Study 2019,” *Lancet*, vol. 396, no. 10258, pp. 1223–1249, Oct. 2020.",
    "India State-Level Disease Burden Initiative Air Pollution Collaborators, “Health and economic impact of air pollution in the states of India: The Global Burden of Disease Study 2019,” *Lancet Planet. Health*, vol. 5, no. 1, pp. e25–e38, Jan. 2021.",
    "P. Orellano, J. Reynoso, N. Quaranta, A. Bardach, and A. Ciapponi, “Short-term exposure to particulate matter (PM10 and PM2.5), nitrogen dioxide (NO2), and ozone (O3) and all-cause and cause-specific mortality: Systematic review and meta-analysis,” *Environ. Int.*, vol. 142, Art. no. 105876, Sep. 2020.",
    "X.-Y. Zheng, H. Ding, L.-N. Jiang, S.-W. Chen, J.-P. Zheng, M. Qiu, Y.-X. Zhou, Q. Chen, and W.-J. Guan, “Association between air pollutants and asthma emergency room visits and hospital admissions in time series studies: A systematic review and meta-analysis,” *PLoS ONE*, vol. 10, no. 9, Art. no. e0138146, Sep. 2015.",
    "Y. Zhang, M. Bocquet, V. Mallet, C. Seigneur, and A. Baklanov, “Real-time air quality forecasting, part I: History, techniques, and current status,” *Atmos. Environ.*, vol. 60, pp. 632–655, Dec. 2012.",
    "X. Li, L. Peng, X. Yao, S. Cui, Y. Hu, C. You, and T. Chi, “Long short-term memory neural network for air pollutant concentration predictions: Method development and evaluation,” *Environ. Pollut.*, vol. 231, pp. 997–1004, Dec. 2017.",
    "S. Ben Taieb, G. Bontempi, A. F. Atiya, and A. Sorjamaa, “A review and comparison of strategies for multi-step ahead time series forecasting based on the NN5 forecasting competition,” *Expert Syst. Appl.*, vol. 39, no. 8, pp. 7067–7083, Jun. 2012.",
    "D. W. Wong, L. Yuan, and S. A. Perlin, “Comparison of spatial interpolation methods for the estimation of air quality data,” *J. Expo. Anal. Environ. Epidemiol.*, vol. 14, no. 5, pp. 404–415, 2004.",
    "C. Chen and B. Zhao, “Review of relationship between indoor and outdoor particles: I/O ratio, infiltration factor and penetration factor,” *Atmos. Environ.*, vol. 45, no. 2, pp. 275–288, Jan. 2011.",
    "K. C. H. Tsang, H. Pinnock, A. M. Wilson, D. Salvi, and S. A. Shah, “Predicting asthma attacks using connected mobile devices and machine learning: The AAMOS-00 observational study protocol,” *BMJ Open*, vol. 12, no. 10, Art. no. e064166, 2022.",
    "S. K. Guttikunda, R. Goel, and P. Pant, “Nature of air pollution, emission sources, and management in the Indian cities,” *Atmos. Environ.*, vol. 95, pp. 501–510, Oct. 2014.",
    "S. Kaufman, S. Rosset, C. Perlich, and O. Stitelman, “Leakage in data mining: Formulation, detection, and avoidance,” *ACM Trans. Knowl. Discov. Data*, vol. 6, no. 4, Art. no. 15, Dec. 2012.",
    "OpenAQ, “OpenAQ API v3 documentation.” [Online]. Available: https://docs.openaq.org (accessed Sep. 2026).",
    "H. Hersbach et al., “The ERA5 global reanalysis,” *Q. J. R. Meteorol. Soc.*, vol. 146, no. 730, pp. 1999–2049, Jul. 2020.",
    "P. Zippenfenig, “Open-Meteo.com weather API.” [Online]. Available: https://open-meteo.com (accessed Sep. 2026).",
    "Central Pollution Control Board, “National Ambient Air Quality Standards,” Gazette of India, Ministry of Environment and Forests, New Delhi, India, Nov. 2009.",
    "G. Ke, Q. Meng, T. Finley, T. Wang, W. Chen, W. Ma, Q. Ye, and T.-Y. Liu, “LightGBM: A highly efficient gradient boosting decision tree,” in *Proc. Adv. Neural Inf. Process. Syst. (NeurIPS)*, Long Beach, CA, USA, 2017, pp. 3146–3154.",
    "S. M. Lundberg and S.-I. Lee, “A unified approach to interpreting model predictions,” in *Proc. Adv. Neural Inf. Process. Syst. (NeurIPS)*, Long Beach, CA, USA, 2017, pp. 4765–4774.",
    "J. C. Platt, “Probabilistic outputs for support vector machines and comparisons to regularized likelihood methods,” in *Advances in Large Margin Classifiers*, A. J. Smola, P. Bartlett, B. Schölkopf, and D. Schuurmans, Eds. Cambridge, MA, USA: MIT Press, 1999, pp. 61–74.",
    "G. W. Brier, “Verification of forecasts expressed in terms of probability,” *Mon. Weather Rev.*, vol. 78, no. 1, pp. 1–3, Jan. 1950.",
    "C. Guo, G. Pleiss, Y. Sun, and K. Q. Weinberger, “On calibration of modern neural networks,” in *Proc. 34th Int. Conf. Mach. Learn. (ICML)*, Sydney, NSW, Australia, 2017, pp. 1321–1330.",
    "S. M. Lundberg, G. Erion, H. Chen, A. DeGrave, J. M. Prutkin, B. Nair, R. Katz, J. Himmelfarb, N. Bansal, and S.-I. Lee, “From local explanations to global understanding with explainable AI for trees,” *Nat. Mach. Intell.*, vol. 2, no. 1, pp. 56–67, Jan. 2020.",
    "S. L. Johnston, P. K. Pattemore, G. Sanderson, S. Smith, F. Lampe, L. Josephs, P. Symington, S. O’Toole, S. H. Myint, D. A. J. Tyrrell, and S. T. Holgate, “Community study of role of viral infections in exacerbations of asthma in 9–11 year old children,” *BMJ*, vol. 310, no. 6989, pp. 1225–1229, May 1995.",
]
for n, ref in enumerate(REFS, start=1):
    reference(n, ref)

props = doc.core_properties
props.title = "AeroGard: Hyperlocal Air-Quality Forecasting, Exposure Estimation and Explainable Respiratory-Risk Alerts"
props.subject = "Research paper"
props.author = "AeroGard project team"
props.keywords = "air quality forecasting, exposure, asthma, COPD, LightGBM, calibration, SHAP, synthetic cohort"
props.comments = ""
props.last_modified_by = ""

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUT))
print(OUT)
